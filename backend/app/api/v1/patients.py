from typing import List, Optional
from uuid import UUID
import io
from fastapi import APIRouter, Depends, HTTPException, status, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.core.database import get_db
from app.models.models import User
from app.api.v1.auth import get_current_user
from app.schemas.patient import PatientCreate, PatientUpdate, PatientResponse
from app.schemas.pagination import PaginatedResponse
from app.crud import patient as patient_crud
from datetime import date, datetime

router = APIRouter()

def require_roles(allowed_roles: List[str]):
    def role_checker(current_user: User = Depends(get_current_user)):
        # Always allow super_admin along with allowed roles
        extended_roles = allowed_roles + ["super_admin"]
        if current_user.role not in extended_roles:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=f"Action forbidden. Required roles: {', '.join(allowed_roles)}"
            )
        return current_user
    return role_checker


@router.post("/", response_model=PatientResponse, status_code=status.HTTP_201_CREATED)
def register_patient(
    patient_in: PatientCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(["admin", "cmo", "nurse", "receptionist"]))
):
    """
    Register a new patient attached to the current user's hospital.
    """
    return patient_crud.create_patient(
        db=db, 
        patient_in=patient_in, 
        current_user_id=current_user.id,
        hospital_id=current_user.hospital_id
    )


@router.get("/", response_model=PaginatedResponse[PatientResponse])
def read_patients(
    search: Optional[str] = None,
    phone: Optional[str] = None,
    registered_from: Optional[date] = Query(None, description="Filter: registered on or after this date"),
    registered_to: Optional[date] = Query(None, description="Filter: registered on or before this date"),
    page: int = 1,
    page_size: int = 20,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(["admin", "doctor", "cmo", "nurse", "receptionist"]))
):
    """
    Search and filter patients scoped to current user's hospital.
    Returns a paginated response with items, total_count, page, and page_size.
    """
    skip = (page - 1) * page_size
    items, total_count = patient_crud.get_patients(
        db=db, 
        search=search, 
        phone=phone,
        registered_from=registered_from,
        registered_to=registered_to, 
        skip=skip, 
        limit=page_size,
        hospital_id=current_user.hospital_id,
        user_role=current_user.role
    )
    return PaginatedResponse(
        items=items,
        total_count=total_count,
        page=page,
        page_size=page_size
    )


@router.get("/export")
def export_patients(
    search: Optional[str] = None,
    phone: Optional[str] = None,
    registered_from: Optional[date] = Query(None, description="Filter: registered on or after this date"),
    registered_to: Optional[date] = Query(None, description="Filter: registered on or before this date"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(["admin", "doctor", "cmo", "nurse", "receptionist"]))
):
    """
    Export the current filtered patient list (matching search/phone filters)
    as an Excel (.xlsx) file.
    """
    patients = patient_crud.get_all_patients_for_export(
        db=db,
        search=search,
        phone=phone,
        registered_from=registered_from,
        registered_to=registered_to,
        hospital_id=current_user.hospital_id,
        user_role=current_user.role
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Patients"

    headers = [
        "First Name", "Last Name", "Date of Birth", "Gender",
        "Phone Number", "Email", "Address",
        "Emergency Contact Name", "Emergency Contact Phone",
        "Blood Group", "Registered On"
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

    for col_num, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_num, patient in enumerate(patients, start=2):
        ws.cell(row=row_num, column=1, value=patient.first_name)
        ws.cell(row=row_num, column=2, value=patient.last_name)
        ws.cell(row=row_num, column=3, value=patient.date_of_birth.strftime("%Y-%m-%d") if patient.date_of_birth else "")
        ws.cell(row=row_num, column=4, value=patient.gender)
        ws.cell(row=row_num, column=5, value=patient.phone_number)
        ws.cell(row=row_num, column=6, value=patient.email or "")
        ws.cell(row=row_num, column=7, value=patient.address)
        ws.cell(row=row_num, column=8, value=patient.emergency_contact_name)
        ws.cell(row=row_num, column=9, value=patient.emergency_contact_phone)
        ws.cell(row=row_num, column=10, value=patient.blood_group or "")
        ws.cell(row=row_num, column=11, value=patient.created_at.strftime("%Y-%m-%d %H:%M") if patient.created_at else "")

    column_widths = [16, 16, 14, 10, 14, 26, 30, 22, 20, 12, 18]
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"patients_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )



@router.get("/{patient_id}", response_model=PatientResponse)
def read_patient_by_id(
    patient_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(["admin", "doctor", "cmo", "nurse", "receptionist"]))
):
    """
    Get details of a single patient by ID (scoped by hospital).
    """
    db_patient = patient_crud.get_patient(
        db=db, 
        patient_id=patient_id,
        hospital_id=current_user.hospital_id,
        user_role=current_user.role
    )
    if not db_patient:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Patient not found"
        )
    return db_patient


@router.put("/{patient_id}", response_model=PatientResponse)
def update_patient_details(
    patient_id: UUID,
    patient_in: PatientUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(["admin", "cmo", "nurse", "receptionist"]))
):
    """
    Update details of an existing patient (scoped by hospital).
    """
    db_patient = patient_crud.get_patient(
        db=db, 
        patient_id=patient_id,
        hospital_id=current_user.hospital_id,
        user_role=current_user.role
    )
    if not db_patient:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Patient not found"
        )
    return patient_crud.update_patient(
        db=db,
        db_obj=db_patient,
        obj_in=patient_in,
        current_user_id=current_user.id
    )


@router.delete("/{patient_id}", status_code=status.HTTP_200_OK)
def remove_patient(
    patient_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(["admin"]))
):
    """
    Soft delete a patient (scoped by hospital).
    """
    db_patient = patient_crud.get_patient(
        db=db, 
        patient_id=patient_id,
        hospital_id=current_user.hospital_id,
        user_role=current_user.role
    )
    if not db_patient:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Patient not found"
        )
    patient_crud.delete_patient(db=db, db_obj=db_patient, current_user_id=current_user.id)
    return {"message": "Patient successfully deleted (soft delete)"}